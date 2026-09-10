"""角色權限矩陣 → Excel（D55）

產出 `docs/角色權限矩陣.xlsx`：哪一頁、哪一個動作，哪種角色可以用。

**矩陣不是手寫的**：頁面與動作那兩欄的勾勾，是拿 `main/utils/permissions.py`
真的跑一次算出來的（假使用者 × has_permission）。程式改了權限、這份表就跟著變，
不會出現「文件說可以、系統說不行」。

用法（在專案根目錄）：
    export POSTGRES_HOST=127.0.0.1 POSTGRES_PORT=30432
    export DJANGO_SETTINGS_MODULE=main.settings.production
    .venv/bin/python shell/export_role_matrix.py

帳號名冊那一頁要連得到資料庫；連不到就只出前三頁（其餘照常）。
"""
import os
import sys

import django

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
os.environ.setdefault("DJANGO_SETTINGS_MODULE", "main.settings.production")
django.setup()

from openpyxl import Workbook  # noqa: E402
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402

from main.apps.core.models import Role  # noqa: E402
from main.utils.permissions import (  # noqa: E402
    NAV_ROLES,
    PERMISSION_ROLES,
    has_permission,
    visible_nav,
)

OUT = os.path.join(
    os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
    "docs", "角色權限矩陣.xlsx",
)

# ── 角色（一個帳號一種）─────────────────────────────────────────────
#   代號 None＝ superuser（admin）。其餘對應 Django Group 的名稱。
ROLES = [
    ("系統管理員", None, "admin",
     "/dashboard", "一人。開帳號、重設密碼、救資料——權限等同經理再加上系統維護"),
    ("經理", Role.OWNER, "manager",
     "/dashboard", "老闆／經理。什麼頁面都看得到、什麼都能改，"
                   "含公司現有現金、產能統計與薪資"),
    ("會計師", Role.FINANCE, "accountant",
     "/dashboard", "會計。什麼頁面都看得到，但只能改金流（應收、應付、登錄付款）、"
                   "算薪水（D57）與自己的任務"),
    ("員工", Role.STAFF, "drafter／clerk／worker01–10",
     "/mywork", "繪圖師、行政人員、工廠員工。除金流外都看得到（唯讀、看不到金額），"
                "能動的是指派給自己的任務與行政事項的完成回報"),
    ("檢視", Role.VIEWER, "（目前沒有人）",
     "/dashboard", "備用角色：同員工的唯讀範圍，但不會被指派任務，也沒有「我的任務」"),
]

ROLE_NAMES = [r[0] for r in ROLES]


class FakeUser:
    """只有角色的假使用者——拿來問 permissions.py「這種角色能不能」"""

    is_authenticated = True

    def __init__(self, code):
        self.is_superuser = code is None
        self.role_codes = {code} if code else set()

    def has_role(self, *codes):
        return self.is_superuser or bool(self.role_codes & set(codes))


USERS = [FakeUser(code) for _, code, *_ in ROLES]

# ── 頁面與動作 ─────────────────────────────────────────────────────
#   (區塊, 項目, 看得到的條件, 能改的權限, 備註)
#   看得到的條件：("nav", 分頁代號) 或 ("perm", 權限代號)
#   能改的權限：None＝這一列只是「看」，有值＝要有那個權限才動得了
PAGE = "nav"
PERM = "perm"

ROWS = [
    ("總覽", "今天要處理什麼（卡片與需要關注）", (PAGE, "dashboard"), None,
     "卡片依角色不同：沒有金流權限的人不會看到收款率、應付這類卡片"),

    ("專案", "案子清單與明細（流程、排程、應收）", (PAGE, "projects"), None,
     "金額欄位另外由「看得到金額」管——員工看得到案子，但金額是空的"),
    ("專案", "新增／編輯／刪除案子、編輯流程", (PAGE, "projects"), "edit_project", ""),
    ("專案", "變更單核准（會動到合約金額）", (PAGE, "projects"), "approve_change_order",
     "牽涉合約金額，只有經理拍板"),

    ("追蹤看板", "流程看板、構件批次、員工視圖、甘特圖", (PAGE, "tracking"), None,
     "跨案看「東西卡在哪一站」，本身不含金額"),
    ("追蹤看板", "改排程日期、狀態、批次過站", (PAGE, "tracking"), "edit_tracking",
     "員工對「指派給自己的單元」可以回報進度——那是資料層的關係，不看角色"),

    ("我的任務", "今天輪到我做什麼（案子＋行政兩種來源）", (PAGE, "mywork"), None,
     "只出現指派給自己的；檢視角色不會被指派，所以沒有這一頁"),

    ("行政", "日曆／條列檢視（全公司的行政事項）", (PAGE, "affairs"), None,
     "所有登入者都看得到全部事項"),
    ("行政", "新增／修改／刪除事項、例行規則、類別", (PAGE, "affairs"), "edit_affairs", ""),
    ("行政", "回報完成、上傳收據與照片", (PAGE, "affairs"), None,
     "被指派的人本人也可以（資料層的關係，不看角色）"),

    ("金流", "應收（錢什麼時候進來）", (PAGE, "finance"), "edit_milestone", ""),
    ("金流", "應付（錢什麼時候出去）", (PAGE, "finance"), "edit_payable", ""),
    ("金流", "應付核可", (PAGE, "finance"), "approve_payable",
     "★ 核可與付款刻意分開：同一個人不該既決定要付多少、又執行付款"),
    ("金流", "登錄付款", (PAGE, "finance"), "pay_payable", ""),
    ("金流", "收支明細（D55：一天一天的收支流水帳）", (PAGE, "finance"), None,
     "看得到金流分頁就看得到；可切「全部／只看案子／只看行政」"),
    ("金流", "現金流預測（哪個月會缺錢）", (PAGE, "finance"), None, ""),
    ("金流", "公司現有現金", (PERM, "view_cash_balance"), "view_cash_balance",
     "連會計師都看不到——這是公司底牌"),

    ("薪資", "月薪資（每個人領多少、每一行怎麼算出來的）", (PAGE, "payroll"), "edit_payroll",
     "★ 比金流更嚴：金流是公司對外的錢，薪資是同事領多少——"
     "所以另立一個權限碼，不跟 view_money 共用"),
    ("薪資", "員工薪資設定（時薪、投保級距、眷屬口數）", (PAGE, "payroll"), "edit_payroll", ""),
    ("薪資", "法規參數（最低工資、加班倍率、勞健保費率）", (PAGE, "payroll"), "edit_payroll",
     "每年修法都會動，所以做成可編輯欄位而不是程式常數"),

    ("統計", "產能與工數（員工做了多少）", (PERM, "view_productivity"), None,
     "涉及員工表現評比，只有經理"),
    ("統計", "金額類統計（單價走勢、流程花費）", (PERM, "view_money"), None, ""),

    ("設定", "客戶、廠商、員工名冊（查電話與聯絡人）", (PAGE, "settings"), None,
     "頁面所有人都看得到，新增／修改的按鈕才跟權限走"),
    ("設定", "維護客戶／廠商／員工、重設密碼", (PAGE, "settings"), "manage_masters", ""),
    ("設定", "流程模板（新案的預設流程；可拖曳排序）", (PERM, "manage_masters"), "manage_masters",
     "這兩個子分頁本身就只有能維護主檔的人看得到，其他角色連分頁都沒有"),
    ("設定", "構件批次站別（幾站、叫什麼、顏色；可拖曳排序）", (PERM, "manage_masters"),
     "manage_masters", "D54 起搬到網站上；Django Admin 已移除"),
]

# ── 權限代號的人話說明（功能權限對照表用）────────────────────────────
PERMISSION_NOTE = {
    "view_overview": "看總覽與專案清單",
    "view_mywork": "有「我的任務」這一頁（會被指派工作）",
    "view_money": "看得到金流分頁與所有金額欄位",
    "view_cash_balance": "看／改公司現有現金",
    "view_productivity": "看產能與工數統計",
    "view_payroll": "看薪資分頁（每位同事的薪水）",
    "edit_payroll": "算薪水、改薪資設定與法規參數",
    "edit_project": "新增、修改、刪除案子與流程",
    "edit_tracking": "改排程、狀態、批次過站",
    "delete_project": "刪除案子",
    "approve_change_order": "核准變更單（會動到合約金額）",
    "edit_milestone": "維護應收款",
    "transition_milestone": "應收款狀態轉換（可請款→已請款→已收款）",
    "edit_subcontract": "維護分包合約",
    "edit_payable": "維護應付款項",
    "approve_payable": "核可應付款項",
    "pay_payable": "登錄付款",
    "manage_masters": "維護主檔（客戶、廠商、員工、流程模板、批次站別）",
    "edit_affairs": "維護行政事項、例行規則與類別",
}

YES, READ, NO = "✔ 可用", "👁 只能看", "✘ 看不到"

# ── 樣式 ───────────────────────────────────────────────────────────
HEAD_FILL = PatternFill("solid", fgColor="1E293B")
HEAD_FONT = Font(color="FFFFFF", bold=True, size=11)
GROUP_FILL = PatternFill("solid", fgColor="F1F5F9")
YES_FILL = PatternFill("solid", fgColor="DCFCE7")
READ_FILL = PatternFill("solid", fgColor="FEF3C7")
NO_FILL = PatternFill("solid", fgColor="FEE2E2")
THIN = Side(style="thin", color="CBD5E1")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER = Alignment(horizontal="center", vertical="center")
WRAP = Alignment(vertical="center", wrap_text=True)

FILL_OF = {YES: YES_FILL, READ: READ_FILL, NO: NO_FILL}


def write_header(ws, titles, widths):
    ws.append(titles)
    for i, (title, width) in enumerate(zip(titles, widths), start=1):
        cell = ws.cell(row=1, column=i)
        cell.fill, cell.font, cell.alignment, cell.border = HEAD_FILL, HEAD_FONT, CENTER, BORDER
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 24


def style_row(ws, row, wrap_cols=()):
    for col in range(1, ws.max_column + 1):
        cell = ws.cell(row=row, column=col)
        cell.border = BORDER
        cell.alignment = WRAP if col in wrap_cols else CENTER
        if cell.value in FILL_OF:
            cell.fill = FILL_OF[cell.value]
            cell.font = Font(bold=True, size=10)


# ── 一、角色定義 ───────────────────────────────────────────────────
def sheet_roles(wb):
    ws = wb.create_sheet("① 角色定義")
    write_header(
        ws,
        ["角色", "代號", "典型帳號", "登入後第一頁", "看得到的分頁", "一句話"],
        [12, 10, 24, 14, 40, 60],
    )
    for name, code, accounts, route, desc in ROLES:
        nav = visible_nav(FakeUser(code))
        ws.append([name, code or "superuser", accounts, route, "、".join(nav), desc])
        style_row(ws, ws.max_row, wrap_cols=(1, 3, 5, 6))
        ws.row_dimensions[ws.max_row].height = 46
    ws.append([])
    ws.append(["規則：一個帳號只掛一種角色。要換角色請到 設定 → 員工 修改，不要一個人掛兩種——"
               "權限取聯集之後，沒有人說得清楚他到底能做什麼。"])
    ws.cell(row=ws.max_row, column=1).font = Font(bold=True, color="B91C1C")
    return ws


# ── 二、頁面權限矩陣 ───────────────────────────────────────────────
def cell_value(user, gate, edit_code):
    kind, key = gate
    if kind == PAGE:
        visible = key in visible_nav(user)
    else:
        visible = has_permission(user, key)
    if not visible:
        return NO
    if edit_code is None:
        return YES        # 這一列本來就只是「看／用」，看得到就是能用
    return YES if has_permission(user, edit_code) else READ


def sheet_matrix(wb):
    ws = wb.create_sheet("② 頁面權限矩陣")
    write_header(
        ws,
        ["分頁", "頁面／動作"] + ROLE_NAMES + ["備註"],
        [12, 42, 13, 13, 13, 13, 13, 52],
    )
    last_group = None
    for group, label, gate, edit_code, note in ROWS:
        values = [cell_value(u, gate, edit_code) for u in USERS]
        ws.append([group if group != last_group else "", label] + values + [note])
        style_row(ws, ws.max_row, wrap_cols=(1, 2, 8))
        if group != last_group:
            for col in (1, 2):
                ws.cell(row=ws.max_row, column=col).fill = GROUP_FILL
            last_group = group
        ws.row_dimensions[ws.max_row].height = 32

    ws.append([])
    ws.append(["圖例", f"{YES}＝可以看也可以改／可以用　　{READ}＝看得到但改不了　　{NO}＝連分頁都不會出現"])
    ws.cell(row=ws.max_row, column=2).font = Font(bold=True)
    ws.append(["", "用不到的分頁是**不顯示**，不是變灰——看到自己用不到的東西只是負荷。"])
    ws.append(["", "員工對「指派給自己的」單元與行政事項，不論角色都能回報進度與勾完成"
                   "（那是資料層的關係，不在這張表裡）。"])
    ws.append(["", "所有金額欄位另外由「看得到金額」把關：員工與檢視在任何頁面拿到的金額都是空的。"])
    return ws


# ── 三、功能權限對照（直接讀程式）──────────────────────────────────
def sheet_permissions(wb):
    ws = wb.create_sheet("③ 功能權限對照")
    write_header(
        ws, ["權限代號", "是什麼"] + ROLE_NAMES, [26, 44, 13, 13, 13, 13, 13],
    )
    for code in PERMISSION_ROLES:
        ws.append(
            [code, PERMISSION_NOTE.get(code, "")]
            + ["✔" if has_permission(u, code) else "✘" for u in USERS]
        )
        style_row(ws, ws.max_row, wrap_cols=(1, 2))
    ws.append([])
    ws.append(["這一頁由 main/utils/permissions.py 直接算出來——程式改了，重跑這支腳本就同步。"])
    ws.cell(row=ws.max_row, column=1).font = Font(italic=True, color="64748B")

    ws.append([])
    ws.append(["分頁", "需要的權限（任一即可，空白＝所有登入者）"])
    for col in (1, 2):
        cell = ws.cell(row=ws.max_row, column=col)
        cell.fill, cell.font = HEAD_FILL, HEAD_FONT
    for nav, required in NAV_ROLES.items():
        ws.append([nav, "、".join(required) if required else "（所有登入者）"])
    return ws


# ── 四、帳號名冊 ───────────────────────────────────────────────────
def sheet_accounts(wb):
    from main.apps.core.models import User

    ws = wb.create_sheet("④ 帳號名冊")
    write_header(ws, ["帳號", "姓名", "職稱", "角色", "狀態"], [16, 16, 16, 14, 12])
    label_of = {code: name for name, code, *_ in ROLES}
    for user in User.objects.all().order_by("employee_no", "username").prefetch_related("groups"):
        codes = [g.name for g in user.groups.all()]
        if user.is_superuser:
            role = "系統管理員"
        elif len(codes) == 1:
            role = label_of.get(codes[0], codes[0])
        elif not codes:
            role = "⚠ 沒有角色"
        else:
            role = "⚠ 掛了多種：" + "、".join(label_of.get(c, c) for c in codes)
        ws.append([user.username, user.name, user.title, role,
                   "啟用中" if user.is_active else "停用"])
        style_row(ws, ws.max_row, wrap_cols=(2, 3, 4))
    ws.append([])
    ws.append(["密碼見 docs/帳號密碼.md。要改某個人的角色：設定 → 員工 → 編輯。"])
    return ws


def main():
    wb = Workbook()
    wb.remove(wb.active)
    sheet_roles(wb)
    sheet_matrix(wb)
    sheet_permissions(wb)
    try:
        sheet_accounts(wb)
    except Exception as exc:   # 連不到資料庫也要出得了前三頁
        print(f"（跳過帳號名冊：{exc}）")
    wb.save(OUT)
    print(f"✔ 已產生 {OUT}")


if __name__ == "__main__":
    main()
