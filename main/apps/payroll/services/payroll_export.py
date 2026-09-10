"""薪資 Excel 匯出（D57）

一個月一個檔，三張表——照會計師接下來真的要做的三件事排：

  ① 薪資總表　　每人一列，出勤與金額全在同一張。給老闆看、自己核對用
  ② 計算明細　　每一行的「項目 × 算式 × 金額」。★ 這張才是重點——
                系統存在的理由就是把算式攤開，匯出去也不能只剩下數字，
                否則印出來的東西跟她原本的 Excel 一樣不能驗算
  ③ 銀行匯款　　編號、姓名、實發金額。她下一步就是拿這個去匯款

刻意不做「一頁四張薪資單」的版面（老闆給的 template 是那樣）：
那是為了列印簽名而存在的格式，而簽名要不要保留還沒確認。
要的話再加第四張表，不影響前面三張。

一個月最多幾十列，直接在記憶體組完即可。
"""
from decimal import Decimal
from io import BytesIO

from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

COLOR_HEAD = "F2F2F2"
COLOR_TOTAL = "E8EEF7"
COLOR_EARNING = "0072B2"
COLOR_DEDUCTION = "D55E00"
COLOR_EMPLOYER = "6E6E6E"

_thin = Side(style="thin", color="BFBFBF")
BORDER = Border(left=_thin, right=_thin, top=_thin, bottom=_thin)

MONEY_FMT = "#,##0"
HOUR_FMT = "0.##"

# 標題, 欄位, 格式, 寬度
SUMMARY_COLUMNS = [
    ("員工編號", "employee_no", None, 11),
    ("姓名", "user_name", None, 12),
    ("計薪方式", "_pay_type", None, 15),
    ("月薪", "_monthly", MONEY_FMT, 10),
    ("上班天數", "work_days", HOUR_FMT, 9),
    ("正常工時", "normal_hours", HOUR_FMT, 9),
    ("平日加班(前段)", "ot_weekday_1_hours", HOUR_FMT, 13),
    ("平日加班(後段)", "ot_weekday_2_hours", HOUR_FMT, 13),
    ("休息日(前2H)", "ot_restday_1_hours", HOUR_FMT, 12),
    ("休息日(3-8H)", "ot_restday_2_hours", HOUR_FMT, 12),
    ("休息日(超過8H)", "ot_restday_3_hours", HOUR_FMT, 13),
    ("國定假日出勤", "holiday_hours", HOUR_FMT, 12),
    ("無薪假", "unpaid_leave_hours", HOUR_FMT, 8),
    ("遲到(分)", "late_minutes", "0", 9),
    ("早退(分)", "early_leave_minutes", "0", 9),
    ("投保薪資", "_insured", MONEY_FMT, 10),
    ("勞保天數", "insured_days", "0", 9),
    ("本月計收健保", "_health", None, 12),
    ("應發總額", "gross", MONEY_FMT, 11),
    ("應扣總額", "deduction", MONEY_FMT, 11),
    ("實發金額", "net", MONEY_FMT, 11),
    ("雇主另負擔", "employer_cost", MONEY_FMT, 11),
]

SECTION_LABEL = {"earning": "應發", "deduction": "應扣", "employer": "雇主負擔"}
SECTION_COLOR = {
    "earning": COLOR_EARNING,
    "deduction": COLOR_DEDUCTION,
    "employer": COLOR_EMPLOYER,
}


def filename(period) -> str:
    """產表日期(YYYYMMDD)＋薪資年月。跟甘特圖匯出同一條命名規則。"""
    ymd = timezone.localdate().strftime("%Y%m%d")
    return f"{ymd}_{period.year}年{period.month:02d}月_薪資表.xlsx"


def _num(value):
    """Decimal 轉成 Excel 認得的數字。存進去是數字才排得了序、加得了總"""
    if value is None:
        return 0
    return float(Decimal(value))


def _title(ws, text, subtitle, span):
    ws.cell(row=1, column=1, value=text).font = Font(bold=True, size=14)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=span)
    ws.cell(row=2, column=1, value=subtitle).font = Font(size=10, color="666666")
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=span)


def _header(ws, row, titles, widths=None):
    fill = PatternFill("solid", fgColor=COLOR_HEAD)
    for i, title in enumerate(titles, start=1):
        cell = ws.cell(row=row, column=i, value=title)
        cell.font = Font(bold=True, size=10)
        cell.fill = fill
        cell.border = BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        if widths:
            ws.column_dimensions[get_column_letter(i)].width = widths[i - 1]
    ws.freeze_panes = ws.cell(row=row + 1, column=1)


def _subtitle(period, records):
    stamp = timezone.localtime().strftime("%Y-%m-%d %H:%M")
    return (
        f"{period.get_status_display()}　"
        f"應上班 {_num(period.workdays):g} 天　"
        f"每日正常工時 {_num(period.normal_hours_per_day):g} 小時　"
        f"共 {len(records)} 人　"
        f"（{stamp} 由系統產生）"
    )


def _summary_sheet(wb, period, records):
    ws = wb.create_sheet("薪資總表")
    _title(ws, f"{period.year} 年 {period.month} 月　薪資總表",
           _subtitle(period, records), len(SUMMARY_COLUMNS))

    head_row = 4
    _header(ws, head_row, [c[0] for c in SUMMARY_COLUMNS], [c[3] for c in SUMMARY_COLUMNS])

    row = head_row
    for record in records:
        row += 1
        for i, (_, field, fmt, _w) in enumerate(SUMMARY_COLUMNS, start=1):
            if field == "employee_no":
                value = record.user.employee_no or ""
            elif field == "user_name":
                value = record.user.name
            elif field == "_pay_type":
                value = _profile_of(record).get_pay_type_display() if _profile_of(record) else "時薪制"
            elif field == "_monthly":
                raw = record.monthly_salary
                if raw is None:
                    profile = _profile_of(record)
                    raw = profile.monthly_salary if profile else None
                value = _num(raw) if raw else ""
            elif field == "_insured":
                value = _num(record.insured_salary if record.insured_salary is not None
                             else _profile_insured(record))
            elif field == "_health":
                value = "計收" if record.charge_health_insurance else "不計收"
            else:
                raw = getattr(record, field)
                value = _num(raw) if fmt else raw
            cell = ws.cell(row=row, column=i, value=value)
            cell.border = BORDER
            if fmt:
                cell.number_format = fmt
                cell.alignment = Alignment(horizontal="right")

    # 合計列。只加總金額欄——把時數加起來沒有意義
    total_row = row + 1
    fill = PatternFill("solid", fgColor=COLOR_TOTAL)
    ws.cell(row=total_row, column=1, value="合計").font = Font(bold=True)
    money_fields = {"gross", "deduction", "net", "employer_cost"}
    for i, (_, field, _fmt, _w) in enumerate(SUMMARY_COLUMNS, start=1):
        cell = ws.cell(row=total_row, column=i)
        cell.fill = fill
        cell.border = BORDER
        cell.font = Font(bold=True)
        if field in money_fields and row > head_row:
            col = get_column_letter(i)
            cell.value = f"=SUM({col}{head_row + 1}:{col}{row})"
            cell.number_format = MONEY_FMT
            cell.alignment = Alignment(horizontal="right")
    return ws


def _profile_of(record):
    return getattr(record.user, "salary_profile", None)


def _profile_insured(record):
    profile = _profile_of(record)
    return profile.insured_salary if profile else 0


def _detail_sheet(wb, period, records):
    """每一行的算式。★ 匯出去也要留著算式，不然就跟原本的 Excel 一樣不能驗算"""
    ws = wb.create_sheet("計算明細")
    titles = ["員工編號", "姓名", "區分", "項目", "算式", "金額"]
    _title(ws, f"{period.year} 年 {period.month} 月　計算明細",
           "每一行都寫著怎麼算出來的——用來核對系統的計算是否正確", len(titles))
    _header(ws, 4, titles, [11, 12, 10, 22, 46, 12])

    row = 4
    for record in records:
        for item in record.detail or []:
            row += 1
            section = item.get("section", "")
            values = [
                record.user.employee_no or "",
                record.user.name,
                SECTION_LABEL.get(section, section),
                item.get("label", ""),
                item.get("formula", ""),
                _num(item.get("amount")),
            ]
            for i, value in enumerate(values, start=1):
                cell = ws.cell(row=row, column=i, value=value)
                cell.border = BORDER
                if i == 3:
                    cell.font = Font(bold=True, size=10, color=SECTION_COLOR.get(section, "000000"))
                if i == 5:
                    # 算式用等寬字，跟畫面上一樣，對數字才對得準
                    cell.font = Font(name="Consolas", size=10)
                if i == 6:
                    cell.number_format = MONEY_FMT
                    cell.alignment = Alignment(horizontal="right")
        # 每個人之間空一列，印出來比較好讀
        row += 1
    return ws


def _bank_sheet(wb, period, records):
    """匯款清單。她核完薪資的下一步就是拿這個去銀行"""
    ws = wb.create_sheet("銀行匯款")
    titles = ["員工編號", "姓名", "實發金額"]
    _title(ws, f"{period.year} 年 {period.month} 月　匯款清單",
           "金額為實發（應發 − 應扣）。帳號不在系統裡，請對照銀行的名冊", len(titles))
    _header(ws, 4, titles, [12, 14, 14])

    row = 4
    for record in records:
        row += 1
        ws.cell(row=row, column=1, value=record.user.employee_no or "").border = BORDER
        ws.cell(row=row, column=2, value=record.user.name).border = BORDER
        cell = ws.cell(row=row, column=3, value=_num(record.net))
        cell.number_format = MONEY_FMT
        cell.border = BORDER
        cell.alignment = Alignment(horizontal="right")

    total_row = row + 1
    ws.cell(row=total_row, column=1, value="合計").font = Font(bold=True)
    cell = ws.cell(row=total_row, column=3)
    if row > 4:
        cell.value = f"=SUM(C5:C{row})"
    cell.number_format = MONEY_FMT
    cell.font = Font(bold=True)
    cell.alignment = Alignment(horizontal="right")
    for col in range(1, 4):
        c = ws.cell(row=total_row, column=col)
        c.fill = PatternFill("solid", fgColor=COLOR_TOTAL)
        c.border = BORDER
    return ws


def build_xlsx(period) -> bytes:
    records = list(
        period.records.select_related("user", "user__salary_profile")
        .order_by("user__employee_no", "user__username")
    )
    wb = Workbook()
    wb.remove(wb.active)          # 預設那張空的 Sheet 不要
    _summary_sheet(wb, period, records)
    _detail_sheet(wb, period, records)
    _bank_sheet(wb, period, records)

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
