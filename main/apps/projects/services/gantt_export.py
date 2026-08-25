"""甘特圖 Excel 匯出

專案頁兩顆下載鈕都走這裡，同一套版面、兩種用途：

  progress（進度追蹤）——內部用。九欄清單（含負責人/狀態/進度）＋圖例，
      顏色＝狀態（灰未開始、淺藍進行中、深藍完成），逾期在狀態欄紅字
  plan（工期規劃）——簽約前給業主看。只留 階段/代號/名稱/起訖/天數，
      沒有圖例、沒有內部欄位，橫條一律深藍（是計畫，不是進度）

共同規則：
  · 檔名與 A1 標題＝產圖日期(YYYYMMDD)＋案名＋用途；A1/A2 合併儲存格
  · 右側日期網格：工期 ≤ 92 天一天一欄，更長一週一欄；月份帶含年份
  · 起訖日期粗體、月/日；跨年的單元補年份（跟網頁同一條規則）
  · 時程網格整片加細框線——貼到別的表、印出來都看得出天數

單案最多 19 列，直接在記憶體組完即可，不需要串流。
"""
from datetime import timedelta
from io import BytesIO

from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from main.utils.choices import FlowState

# 與網頁一致的狀態色（Okabe–Ito 藍系）
COLOR_TODO = "D9D9D9"
COLOR_DOING = "56B4E9"
COLOR_DONE = "0072B2"
COLOR_RED = "D55E00"
COLOR_HEAD = "F2F2F2"

INFO_PROGRESS = ["階段", "代號", "流程名稱", "負責人", "狀態", "預計開始", "預計完成", "天數", "進度"]
INFO_PLAN = ["階段", "代號", "流程名稱", "預計開始", "預計完成", "天數"]
WIDTHS_PROGRESS = [13, 5.5, 24, 10, 13, 9, 9, 6, 12]
WIDTHS_PLAN = [13, 5.5, 24, 9, 9, 6]

# 一天一欄的上限。超過就一週一欄，不然半年的案子會有 180 欄
DAY_MODE_MAX_DAYS = 92

_thin = Side(style="thin", color="BFBFBF")
GRID_BORDER = Border(left=_thin, right=_thin, top=_thin, bottom=_thin)


def variant_label(variant: str) -> str:
    return "工期規劃" if variant == "plan" else "進度追蹤"


def filename(project, variant: str) -> str:
    """產圖日期(YYYYMMDD)＋案名＋用途。案名裡的路徑符號換掉，檔名才安全。"""
    ymd = timezone.localdate().strftime("%Y%m%d")
    safe_name = project.name.replace("/", "-").replace("\\", "-")
    return f"{ymd}_{safe_name}_{variant_label(variant)}.xlsx"


def build_xlsx(project, variant: str = "progress") -> bytes:
    plan = variant == "plan"
    headers = INFO_PLAN if plan else INFO_PROGRESS
    grid_first = len(headers) + 1
    # 版面：1 標題、2 產表資訊、（進度版 3 圖例）、月份帶、日期列、資料列
    month_row = 3 if plan else 4
    header_row = month_row + 1
    first_data = header_row + 1

    units = list(
        project.flow_units.select_related("flow_item__stage", "assignee")
        .prefetch_related("tasks__assignments")  # 進度由工作分配算（D44）
        .exclude(state=FlowState.NA)
        .order_by("flow_item__seq")
    )
    today = timezone.localdate()

    wb = Workbook()
    ws = wb.active
    ws.title = variant_label(variant)

    dated = [u for u in units if u.plan_start and u.plan_end]
    span_start = min((u.plan_start for u in dated), default=None)
    span_end = max((u.plan_end for u in dated), default=None)

    _write_title_rows(ws, project, variant, today, span_start, span_end, len(headers))
    if not plan:
        _write_legend(ws)

    for idx, name in enumerate(headers, start=1):
        cell = ws.cell(row=header_row, column=idx, value=name)
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor=COLOR_HEAD)

    columns = _build_grid_columns(ws, span_start, span_end, today, grid_first, month_row, header_row)

    row = first_data
    for unit in units:
        _write_unit_row(ws, row, unit, columns, today, plan)
        row += 1
    if not units:
        ws.cell(row=row, column=1, value="這個案子還沒勾選流程")

    # 時程網格整片細框線（含月份帶與日期列）——沒排日期就沒有網格可框
    if columns:
        last_col = columns[-1][2]
        for r in range(month_row, first_data + max(len(units), 1)):
            for c in range(grid_first, last_col + 1):
                ws.cell(row=r, column=c).border = GRID_BORDER

    widths = WIDTHS_PLAN if plan else WIDTHS_PROGRESS
    for idx, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width
    day_mode = bool(columns) and columns[0][0] == columns[0][1]
    for _, _, col in columns:
        ws.column_dimensions[get_column_letter(col)].width = 3.4 if day_mode else 6

    ws.freeze_panes = f"{get_column_letter(grid_first)}{first_data}"  # 捲動時表頭與左欄不動

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ── 標題與產表資訊：A1/A2 各自合併到資訊欄最右 ─────────────────────
def _write_title_rows(ws, project, variant, today, span_start, span_end, info_cols):
    title = ws.cell(row=1, column=1, value=f"{today.strftime('%Y%m%d')} {project.name} {variant_label(variant)}")
    title.font = Font(bold=True, size=14)
    title.alignment = Alignment(horizontal="left", vertical="center")

    span_text = (
        f"時間軸 {span_start.strftime('%Y/%m/%d')} ~ {span_end.strftime('%Y/%m/%d')}"
        if span_start
        else "尚無排程日期"
    )
    meta = ws.cell(row=2, column=1, value=f"產表日期 {today.strftime('%Y/%m/%d')}　{span_text}")
    meta.alignment = Alignment(horizontal="left", vertical="center")

    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=info_cols)
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=info_cols)


def _write_legend(ws):
    ws.cell(row=3, column=1, value="圖例")
    col = 2
    for color, label in [(COLOR_TODO, "未開始"), (COLOR_DOING, "進行中"), (COLOR_DONE, "已完成")]:
        ws.cell(row=3, column=col).fill = PatternFill("solid", fgColor=color)
        ws.cell(row=3, column=col + 1, value=label)
        col += 2
    note = ws.cell(row=3, column=col, value="逾期＝狀態欄紅字（不單靠顏色）")
    note.font = Font(color="808080", size=10)


# ── 日期網格的欄：回傳 [(起日, 迄日, 欄號), ...] ───────────────────
def _build_grid_columns(ws, span_start, span_end, today, grid_first, month_row, header_row):
    if not span_start:
        return []

    day_mode = (span_end - span_start).days + 1 <= DAY_MODE_MAX_DAYS
    columns = []
    col = grid_first

    if day_mode:
        d = span_start
        while d <= span_end:
            columns.append((d, d, col))
            cell = ws.cell(row=header_row, column=col, value=d.day)
            cell.alignment = Alignment(horizontal="center")
            cell.font = Font(
                size=9,
                bold=(d == today),
                color=COLOR_RED if d == today else ("808080" if d.weekday() >= 5 else "000000"),
            )
            col += 1
            d += timedelta(days=1)
    else:
        # 對齊到週一，一欄一週，欄頭標該週週一的 月/日
        d = span_start - timedelta(days=span_start.weekday())
        while d <= span_end:
            week_end = d + timedelta(days=6)
            columns.append((d, week_end, col))
            cell = ws.cell(row=header_row, column=col, value=f"{d.month}/{d.day}")
            cell.alignment = Alignment(horizontal="center")
            cell.font = Font(size=9, bold=d <= today <= week_end,
                             color=COLOR_RED if d <= today <= week_end else "000000")
            col += 1
            d += timedelta(days=7)

    # 月份帶（合併儲存格，標「YYYY年M月」——跨年也看得懂）
    month_start_idx = 0
    for i in range(1, len(columns) + 1):
        boundary = i == len(columns) or (
            columns[i][0].month != columns[month_start_idx][0].month
            or columns[i][0].year != columns[month_start_idx][0].year
        )
        if not boundary:
            continue
        first = columns[month_start_idx]
        last = columns[i - 1]
        cell = ws.cell(row=month_row, column=first[2], value=f"{first[0].year}年{first[0].month}月")
        cell.font = Font(bold=True, size=10)
        cell.fill = PatternFill("solid", fgColor=COLOR_HEAD)
        if last[2] > first[2]:
            ws.merge_cells(start_row=month_row, start_column=first[2], end_row=month_row, end_column=last[2])
        month_start_idx = i

    return columns


# ── 一列＝一張流程單元 ─────────────────────────────────────────────
def _write_unit_row(ws, row, unit, columns, today, plan):
    item = unit.flow_item
    ws.cell(row=row, column=1, value=f"{item.stage.seq} {item.stage.name}")
    ws.cell(row=row, column=2, value=item.code)
    ws.cell(row=row, column=3, value=item.name)

    if plan:
        date_cols = (4, 5)
        days_col = 6
    else:
        ws.cell(row=row, column=4, value=unit.assignee.name if unit.assignee else "")
        overdue = (
            unit.state in (FlowState.TODO, FlowState.DOING)
            and unit.plan_end is not None
            and unit.plan_end < today
        )
        state_cell = ws.cell(
            row=row, column=5,
            value=f"{unit.get_state_display()}（已逾期）" if overdue else unit.get_state_display(),
        )
        if overdue:
            state_cell.font = Font(bold=True, color=COLOR_RED)
        ws.cell(row=row, column=9, value=_progress_text(unit))
        date_cols = (6, 7)
        days_col = 8

    # 起訖：粗體、顯示 月/日；單元跨年時補年份。儲存格放真正的日期值，
    # 之後要拿去算天數、排序都不會掉年份
    if unit.plan_start and unit.plan_end:
        cross_year = unit.plan_start.year != unit.plan_end.year
        fmt = "yyyy/m/d" if cross_year else "m/d"
        for col, value in zip(date_cols, (unit.plan_start, unit.plan_end), strict=False):
            cell = ws.cell(row=row, column=col, value=value)
            cell.number_format = fmt
            cell.font = Font(bold=True)
        ws.cell(row=row, column=days_col, value=(unit.plan_end - unit.plan_start).days + 1)
    else:
        return

    # 橫條：業主版一律深藍（計畫）；進度版照狀態著色
    fill_color = COLOR_DONE if plan else {
        FlowState.TODO: COLOR_TODO,
        FlowState.DOING: COLOR_DOING,
        FlowState.DONE: COLOR_DONE,
    }[unit.state]
    fill = PatternFill("solid", fgColor=fill_color)
    for col_start, col_end, col in columns:
        if col_end >= unit.plan_start and col_start <= unit.plan_end:
            ws.cell(row=row, column=col).fill = fill


def _progress_text(unit):
    if unit.qty_total:
        qty = f"{_num(unit.qty_done)}/{_num(unit.qty_total)}"
        return f"{qty} {unit.unit_of_measure}".strip()
    return f"{_num(unit.completion_ratio)}%"


def _num(value):
    f = float(value)
    return str(int(f)) if f == int(f) else f"{f:g}"
